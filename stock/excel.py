"""Import et export Excel : stock (aller-retour au même format) et ventes."""

import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services
from .models import Ajustement, Categorie, Reference, Vente

# Colonnes du fichier de stock, dans l'ordre de l'export. L'import les retrouve par leur
# titre (sans tenir compte des accents ni des majuscules), où qu'elles soient dans la feuille.
COLONNES_STOCK = ["DÉSIGNATION", "CATÉGORIE", "COÛT D'ACHAT", "PRIX DE VENTE", "QUANTITÉ EN STOCK", "ACTIVE"]
OBLIGATOIRES = COLONNES_STOCK[:5]
SANS_CATEGORIE = {"", "-", "—"}


def _cle(texte):
    texte = unicodedata.normalize("NFKD", str(texte or "")).encode("ascii", "ignore").decode()
    return " ".join(texte.replace("’", "'").upper().split())


def _montant(valeur):
    if valeur is None or str(valeur).strip() == "":
        raise ValueError("vide")
    try:
        montant = Decimal(str(valeur).strip().replace(",", ".").replace("€", "").replace(" ", ""))
    except InvalidOperation:
        raise ValueError(f"« {valeur} » n'est pas un nombre")
    if montant < 0:
        raise ValueError("négatif")
    return montant.quantize(Decimal("0.01"))


def _quantite(valeur):
    montant = _montant(valeur)
    if montant != montant.to_integral_value():
        raise ValueError(f"« {valeur} » n'est pas un nombre entier")
    return int(montant)


def _booleen(valeur):
    if valeur is None or str(valeur).strip() == "":
        return True
    return _cle(valeur) in {"TRUE", "VRAI", "OUI", "1", "X", "ACTIVE", "ACTIF"}


@dataclass
class LigneImport:
    numero: int
    designation: str
    categorie: str
    cout_achat: Decimal
    prix_vente: Decimal
    quantite: int
    active: bool
    action: str = ""  # « création » ou « mise à jour » ou « inchangée »
    details: list = field(default_factory=list)

    def vers_session(self):
        return {**self.__dict__, "cout_achat": str(self.cout_achat), "prix_vente": str(self.prix_vente)}

    @classmethod
    def depuis_session(cls, d):
        return cls(**{**d, "cout_achat": Decimal(d["cout_achat"]), "prix_vente": Decimal(d["prix_vente"])})


@dataclass
class AnalyseImport:
    lignes: list = field(default_factory=list)
    erreurs: list = field(default_factory=list)
    avertissements: list = field(default_factory=list)
    nouvelles_categories: list = field(default_factory=list)


def analyser_stock(fichier):
    """Lit le fichier et prépare l'import, sans rien enregistrer."""
    analyse = AnalyseImport()
    try:
        feuille = load_workbook(fichier, data_only=True, read_only=True).worksheets[0]
    except Exception:
        analyse.erreurs.append("Le fichier n'est pas un fichier Excel (.xlsx) lisible.")
        return analyse

    lignes = feuille.iter_rows(values_only=True)
    entetes = None
    for numero, ligne in enumerate(lignes, start=1):
        cles = [_cle(v) for v in ligne]
        if _cle("DÉSIGNATION") in cles:
            entetes = {cle: index for index, cle in enumerate(cles) if cle}
            break
        if numero >= 10:
            break
    if entetes is None:
        analyse.erreurs.append("Ligne de titres introuvable : la feuille doit contenir une colonne « DÉSIGNATION ».")
        return analyse
    manquantes = [c for c in OBLIGATOIRES if _cle(c) not in entetes]
    if manquantes:
        analyse.erreurs.append("Colonnes manquantes : " + ", ".join(manquantes) + ".")
        return analyse

    def cellule(ligne, colonne):
        index = entetes.get(_cle(colonne))
        return ligne[index] if index is not None and index < len(ligne) else None

    existantes = {_cle(r.designation): r for r in Reference.objects.select_related("categorie")}
    categories = {_cle(c.nom): c.nom for c in Categorie.objects.all()}
    vues = set()
    for numero, ligne in enumerate(lignes, start=numero + 1):
        designation = str(cellule(ligne, "DÉSIGNATION") or "").strip()
        if not designation:
            continue  # ligne vide ou ligne de total
        try:
            cout = _montant(cellule(ligne, "COÛT D'ACHAT"))
            prix = _montant(cellule(ligne, "PRIX DE VENTE"))
            quantite = _quantite(cellule(ligne, "QUANTITÉ EN STOCK"))
        except ValueError as e:
            analyse.erreurs.append(f"Ligne {numero} ({designation}) : valeur invalide — {e}.")
            continue
        if _cle(designation) in vues:
            analyse.erreurs.append(f"Ligne {numero} : « {designation} » apparaît plusieurs fois dans le fichier.")
            continue
        vues.add(_cle(designation))

        categorie = str(cellule(ligne, "CATÉGORIE") or "").strip()
        if categorie in SANS_CATEGORIE:
            categorie = ""
            analyse.avertissements.append(f"Ligne {numero} ({designation}) : sans catégorie.")
        elif _cle(categorie) in categories:
            categorie = categories[_cle(categorie)]
        elif categorie not in analyse.nouvelles_categories:
            analyse.nouvelles_categories.append(categorie)
        if prix < cout:
            analyse.avertissements.append(f"Ligne {numero} ({designation}) : prix de vente inférieur au coût d'achat.")

        l = LigneImport(numero, designation, categorie, cout, prix, quantite, _booleen(cellule(ligne, "ACTIVE")))
        ref = existantes.get(_cle(designation))
        if ref is None:
            l.action = "création"
        else:
            l.designation = ref.designation
            changements = []
            if (ref.categorie.nom if ref.categorie else "") != categorie:
                changements.append(f"catégorie {ref.categorie or '—'} → {categorie or '—'}")
            if ref.cout_achat != cout:
                changements.append(f"coût {ref.cout_achat} → {cout} €")
            if ref.prix_vente != prix:
                changements.append(f"prix {ref.prix_vente} → {prix} €")
            if ref.quantite != quantite:
                changements.append(f"stock {ref.quantite} → {quantite}")
            if ref.active != l.active:
                changements.append("réactivée" if l.active else "masquée")
            l.action = "mise à jour" if changements else "inchangée"
            l.details = changements
        analyse.lignes.append(l)
    return analyse


@transaction.atomic
def importer_stock(lignes, utilisateur):
    """Enregistre les lignes analysées. Les quantités passent par un ajustement « Import Excel »
    pour garder la trace des mouvements de stock."""
    compte = {"création": 0, "mise à jour": 0, "inchangée": 0}
    categories = {}
    for l in lignes:
        categorie = None
        if l.categorie:
            if l.categorie not in categories:
                categories[l.categorie], _ = Categorie.objects.get_or_create(nom=l.categorie)
            categorie = categories[l.categorie]
        ref = Reference.objects.filter(designation__iexact=l.designation).first()
        if ref is None:
            ref = Reference.objects.create(
                designation=l.designation, categorie=categorie, cout_achat=l.cout_achat,
                prix_vente=l.prix_vente, quantite=0, active=l.active,
            )
            ecart = l.quantite
        else:
            ref.categorie, ref.cout_achat, ref.prix_vente, ref.active = categorie, l.cout_achat, l.prix_vente, l.active
            ref.save()
            ecart = l.quantite - ref.quantite
        if ecart:
            services.ajuster_stock(ref.pk, ecart, Ajustement.Motif.INVENTAIRE, utilisateur, "Import Excel")
        compte[l.action] = compte.get(l.action, 0) + 1
    return compte


def _classeur(titre_feuille, colonnes, lignes, largeurs):
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = titre_feuille
    feuille.append(colonnes)
    for cellule in feuille[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="C2185B")
    for ligne in lignes:
        feuille.append(ligne)
    for i, largeur in enumerate(largeurs, start=1):
        feuille.column_dimensions[get_column_letter(i)].width = largeur
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions
    return classeur


def exporter_stock():
    refs = Reference.objects.select_related("categorie").order_by("categorie__nom", "designation")
    lignes = [
        [r.designation, r.categorie.nom if r.categorie else "", float(r.cout_achat), float(r.prix_vente),
         r.quantite, "Oui" if r.active else "Non", float(r.cout_achat * r.quantite), float(r.prix_vente * r.quantite)]
        for r in refs
    ]
    classeur = _classeur("Stock", COLONNES_STOCK + ["VALEUR D'ACHAT", "VALEUR DE VENTE"], lignes,
                         [42, 22, 14, 14, 18, 10, 16, 16])
    feuille = classeur.active
    for colonne in "CDGH":
        for cellule in feuille[colonne][1:]:
            cellule.number_format = '#,##0.00 "€"'
    return classeur


def exporter_ventes(debut, fin):
    ventes = (
        Vente.objects.filter(date__date__gte=debut, date__date__lte=fin)
        .select_related("reference__categorie", "vendeur", "saisie_par").order_by("date")
    )
    lignes = []
    for v in ventes:
        moment = timezone.localtime(v.date)
        lignes.append([
            moment.date(), moment.strftime("%H:%M"), v.reference.designation,
            v.reference.categorie.nom if v.reference.categorie else "", v.quantite,
            float(v.prix_catalogue), v.remise_pct or 0, float(v.prix_unitaire), float(v.chiffre_affaires),
            float(v.cout_unitaire * v.quantite), float(v.marge), v.get_canal_display(),
            v.get_paiement_display() if v.paiement else "", float(v.montant_especes), float(v.montant_cb),
            (v.vendeur.get_short_name() or v.vendeur.username) if v.vendeur else "",
            (v.saisie_par.get_short_name() or v.saisie_par.username) if v.saisie_par else "",
        ])
    colonnes = ["DATE", "HEURE", "RÉFÉRENCE", "CATÉGORIE", "QUANTITÉ", "PRIX INITIAL", "SOLDE (%)",
                "PRIX UNITAIRE", "TOTAL", "COÛT D'ACHAT", "MARGE", "CANAL", "PAIEMENT", "ESPÈCES", "CB",
                "VENDEUR", "SAISIE PAR"]
    classeur = _classeur("Ventes", colonnes, lignes, [12, 8, 40, 20, 10, 13, 10, 14, 12, 14, 12, 11, 15, 11, 11, 14, 14])
    feuille = classeur.active
    for cellule in feuille["A"][1:]:
        cellule.number_format = "DD/MM/YYYY"
    for colonne in "FHIJKNO":
        for cellule in feuille[colonne][1:]:
            cellule.number_format = '#,##0.00 "€"'
    return classeur
