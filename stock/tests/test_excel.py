import io
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from stock.models import Ajustement, Categorie, Reference, User
from stock.tests.outils import vendre

TITRES = [None, "DÉSIGNATION", "CATÉGORIE", "COÛT D'ACHAT", "PRIX DE VENTE", "QUANTITÉ EN STOCK", "ACTIVE", None, None]


def fichier(lignes, titres=TITRES):
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(titres)
    for ligne in lignes:
        feuille.append(ligne)
    flux = io.BytesIO()
    classeur.save(flux)
    flux.seek(0)
    flux.name = "Stock.xlsx"
    return flux


class ImportStockTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user("garnik", password="x", role=User.Role.ADMIN)
        cls.lea = User.objects.create_user("lea", password="x")

    def setUp(self):
        self.client.force_login(self.admin)

    def _verifier(self, lignes, **kw):
        return self.client.post(reverse("excel_import_stock"), {"fichier": fichier(lignes, **kw)})

    def _confirmer(self):
        return self.client.post(reverse("excel_import_stock"), {"confirmer": "1"}, follow=True)

    def test_apercu_puis_import(self):
        r = self._verifier([
            [None, "Body dentelle noir", "Body", 10, 39, 2, "True", None, 20],
            [None, "Chemise à pois ecru", "Chemise", 15.5, 29, 1, "True", None, 9],
            [None, "Haut dentelle taupe", "-", 5, 15, 1, "True", None, 5],
            [None, None, None, None, None, None, None, None, 34],  # ligne de total : ignorée
        ])
        self.assertContains(r, "3</strong> références lues")
        self.assertContains(r, "Nouvelles catégories créées : <strong>Body</strong>")
        self.assertContains(r, "sans catégorie")
        self.assertFalse(Reference.objects.exists())  # rien avant confirmation

        r = self._confirmer()
        self.assertContains(r, "3 référence(s) créée(s)")
        body = Reference.objects.get(designation="Body dentelle noir")
        self.assertEqual((body.categorie.nom, body.cout_achat, body.prix_vente, body.quantite, body.active),
                         ("Body", Decimal("10.00"), Decimal("39.00"), 2, True))
        self.assertEqual(Reference.objects.get(designation="Chemise à pois ecru").cout_achat, Decimal("15.50"))
        self.assertIsNone(Reference.objects.get(designation="Haut dentelle taupe").categorie)
        self.assertEqual(Ajustement.objects.filter(commentaire="Import Excel").count(), 3)

    def test_mise_a_jour_d_une_reference_existante(self):
        pull = Categorie.objects.get(nom="Pull")
        Reference.objects.create(designation="Pull angora", categorie=pull, cout_achat=20, prix_vente=50, quantite=5)
        r = self._verifier([[None, "pull ANGORA", "Pull", 22, 55, 3, "True"]])
        self.assertContains(r, "stock 5 → 3")
        self._confirmer()
        ref = Reference.objects.get()
        self.assertEqual((ref.cout_achat, ref.prix_vente, ref.quantite), (Decimal("22.00"), Decimal("55.00"), 3))
        ajustement = Ajustement.objects.get()
        self.assertEqual((ajustement.variation, ajustement.motif), (-2, Ajustement.Motif.INVENTAIRE))

    def test_erreurs_bloquantes(self):
        r = self._verifier([[None, "Pull", "Pull", "abc", 55, 3, "True"], [None, "Jupe", "Jupe", 10, 30, 1.5, "True"]])
        self.assertContains(r, "Rien n'a été importé")
        self.assertContains(r, "n&#x27;est pas un nombre")
        self.assertContains(r, "n&#x27;est pas un nombre entier")
        self.assertNotContains(r, "Confirmer l'import")
        r = self._confirmer()
        self.assertContains(r, "Aucun import en attente")
        self.assertFalse(Reference.objects.exists())

    def test_colonnes_manquantes_et_doublons(self):
        r = self._verifier([[None, "Pull"]], titres=[None, "DÉSIGNATION", "CATÉGORIE"])
        self.assertContains(r, "Colonnes manquantes")
        r = self._verifier([[None, "Pull", "Pull", 1, 2, 1, "True"], [None, "pull", "Pull", 1, 2, 1, "True"]])
        self.assertContains(r, "plusieurs fois")

    def test_fichier_illisible(self):
        flux = io.BytesIO(b"pas un fichier excel")
        flux.name = "x.xlsx"
        r = self.client.post(reverse("excel_import_stock"), {"fichier": flux})
        self.assertContains(r, "pas un fichier Excel")

    def test_reserve_aux_administrateurs(self):
        self.client.force_login(self.lea)
        for url in ("excel_accueil", "excel_export_stock"):
            self.assertNotEqual(self.client.get(reverse(url)).status_code, 200)
        self.client.post(reverse("excel_import_stock"), {"fichier": fichier([[None, "Pull", "Pull", 1, 2, 1]])})
        self.client.post(reverse("excel_import_stock"), {"confirmer": "1"})
        self.assertFalse(Reference.objects.exists())


class ExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = User.objects.create_user("garnik", password="x", role=User.Role.ADMIN)
        cls.lea = User.objects.create_user("lea", password="x", first_name="Léa")
        cls.robe = Reference.objects.create(designation="Robe lin", categorie=Categorie.objects.get(nom="Pull"),
                                            cout_achat=Decimal("30"), prix_vente=Decimal("50"), quantite=10)

    def _classeur(self, reponse):
        self.assertEqual(reponse.status_code, 200)
        self.assertIn("spreadsheetml", reponse["Content-Type"])
        return load_workbook(io.BytesIO(reponse.content)).active

    def test_export_stock_reimportable(self):
        self.client.force_login(self.admin)
        feuille = self._classeur(self.client.get(reverse("excel_export_stock")))
        lignes = list(feuille.iter_rows(values_only=True))
        self.assertEqual(lignes[0][:6], ("DÉSIGNATION", "CATÉGORIE", "COÛT D'ACHAT", "PRIX DE VENTE",
                                         "QUANTITÉ EN STOCK", "ACTIVE"))
        self.assertEqual(lignes[1][:8], ("Robe lin", "Pull", 30, 50, 10, "Oui", 300, 500))
        # Le fichier exporté se réimporte sans changement.
        flux = io.BytesIO(self.client.get(reverse("excel_export_stock")).content)
        flux.name = "stock.xlsx"
        r = self.client.post(reverse("excel_import_stock"), {"fichier": flux})
        self.assertContains(r, "inchangée")
        self.assertEqual(r.context["mises_a_jour"], [])

    def test_export_ventes(self):
        vendre(self.robe.pk, 2, self.lea, remise_pct=10, paiement="mixte", montant_especes=Decimal("40"))
        self.client.force_login(self.admin)
        r = self.client.get(reverse("excel_export_ventes"), {"debut": "2000-01-01", "fin": "2100-01-01"})
        lignes = list(self._classeur(r).iter_rows(values_only=True))
        self.assertEqual(len(lignes), 2)
        entetes = dict(zip(lignes[0], lignes[1]))
        self.assertEqual(entetes["RÉFÉRENCE"], "Robe lin")
        self.assertEqual(entetes["TOTAL"], 90)
        self.assertEqual(entetes["SOLDE (%)"], 10)
        self.assertEqual(entetes["ESPÈCES"], 40)
        self.assertEqual(entetes["CB"], 50)
        self.assertEqual(entetes["VENDEUR"], "Léa")

    def test_lien_dans_l_administration(self):
        self.client.force_login(User.objects.create_superuser("chef", "", "x"))
        self.assertContains(self.client.get(reverse("admin:index")), reverse("excel_accueil"))
